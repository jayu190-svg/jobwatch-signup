"""Job Watch sign-up page (Streamlit).

A friend enters an invite code, fills in a short form, uploads a resume (PDF/DOCX), and this app commits
profiles/<id>/profile.json + resume.txt to the PRIVATE repo 'jobwatch-profiles'. The Job Watch engine on apex-vm
pulls that repo at the start of every run, so the person appears automatically.

Secrets (Streamlit "Secrets", never in code):  INVITE_CODE, GITHUB_TOKEN (repo scope).  Optional: PROFILES_REPO ("owner/name").
Everything above the `main()` line is plain functions so it can be tested without a browser.
"""
import base64
import hashlib
import hmac
import io
import json
import re
import secrets as pysecrets
from datetime import datetime, timezone

import requests

API = "https://api.github.com"
PROFILES_REPO_NAME = "jobwatch-profiles"
MAX_TITLES, MAX_CITIES, MAX_CORE = 25, 5, 60          # 25 titles x 5 cities = the engine's 125-call ceiling
MAX_RESUME_BYTES, MAX_RESUME_CHARS = 5_000_000, 20_000
LEVELS = {"entry_mid": "Entry to mid level", "any": "Any level", "senior": "Senior"}
DEFAULT_EXCLUDE = ["security clearance", "top secret", "ts/sci"]       # same defaults the engine's template uses
DEFAULT_AVOID_TITLE = ["intern", "sales", "recruiter"]
PUSHOVER_KEY_RE = re.compile(r"^[A-Za-z0-9]{20,40}$")
EMAIL_RE = re.compile(r"^[^@\s]{1,64}@[^@\s]{1,255}\.[^@\s]{2,}$")
GIST_PREFIX = "https://gist.github.com/"


# ---------------------------------------------------------------- pure helpers
def parse_items(text: str, max_items: int, max_len: int, sep: str = r"[\n;]") -> list:
    """Split free text into a de-duplicated list (case-insensitive), dropping blanks and control characters."""
    out, seen = [], set()
    for part in re.split(sep, text or ""):
        s = re.sub(r"[\x00-\x1f\x7f]", " ", part).strip()[:max_len]
        if s and s.lower() not in seen:
            seen.add(s.lower())
            out.append(s)
    return out[:max_items]


def slugify(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:30] or "friend"


def resume_to_text(filename: str, data: bytes) -> str:
    """PDF (pypdf) or DOCX (python-docx) -> plain text. Raises ValueError with a message that is safe to show."""
    if len(data) > MAX_RESUME_BYTES:
        raise ValueError("That file is larger than 5 MB. Please upload a smaller resume.")
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("That PDF is password-protected. Please upload an unprotected copy.")
            text = "\n".join((pg.extract_text() or "") for pg in reader.pages[:10])
        elif name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                parts += [c.text for row in t.rows for c in row.cells]
            text = "\n".join(parts)
        else:
            raise ValueError("Please upload a PDF or a Word (.docx) file.")
    except ValueError:
        raise
    except Exception:
        raise ValueError("I couldn't read that file. Try saving it again as a PDF or .docx.")
    text = re.sub(r"[ \t]+", " ", re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", text))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if len(text) < 200:
        raise ValueError("I could read almost no text in that file (is it a scanned image?). Please upload a text-based PDF or .docx.")
    return text[:MAX_RESUME_CHARS]


def validate_form(f: dict, resume_required: bool = True) -> list:
    errs = []
    if not f["first_name"]:
        errs.append("First name is required.")
    if f.get("email") is not None and f["email"] and not EMAIL_RE.match(f["email"]):
        errs.append("That email address doesn't look right.")
    if not f["titles"]:
        errs.append("Enter at least one job title.")
    if not f["core_words"]:
        errs.append("Enter at least one core word.")
    if not f["remote_only"] and not f["cities"]:
        errs.append("Enter at least one city, or tick 'Remote only'.")
    if len(f["cities"]) > MAX_CITIES:
        errs.append(f"Please list at most {MAX_CITIES} cities.")
    if f["pushover_user"] and not PUSHOVER_KEY_RE.match(f["pushover_user"]):
        errs.append("The Pushover user key should be 20-40 letters and numbers (leave it blank if you don't use Pushover).")
    if resume_required and not f.get("resume_text"):
        errs.append("Please upload your resume.")
    return errs


def build_profile(f: dict) -> dict:
    """Form answers -> profile.json in the engine's schema. The engine re-sanitizes it, so this only has to be well-formed.
    '_form' keeps the raw answers so the Edit tab can pre-fill the form; the engine ignores it."""
    cities = ["Remote"] if f["remote_only"] else f["cities"]
    where = "remote only" if f["remote_only"] else "; ".join(f["cities"])
    return {
        "first_name": f["first_name"],
        "name": f["first_name"],
        "summary": (f"{f['first_name']} is looking for roles like: {', '.join(f['titles'][:8])}. Core skills and words: {', '.join(f['core_words'][:15])}. "
                    f"Location: {where}. Level: {LEVELS[f['level']].lower()}.")[:1500],
        "searches": f["titles"][:MAX_TITLES],
        "tracks": {"Main": f["core_words"][:MAX_CORE]},
        "anchor_tracks": ["Main"],
        "locations": cities,
        "location_keywords": [] if f["remote_only"] else [c.split(",")[0].strip().lower() for c in f["cities"] if c.lower() != "remote"],
        "remote_only": bool(f["remote_only"]),
        "level": {"prefer": f["level"]},
        "min_salary": int(f["min_salary"]),
        "exclude_if_required": DEFAULT_EXCLUDE + [x for x in f["lacks"] if x.lower() not in DEFAULT_EXCLUDE],
        "blocked_companies": f["avoid"],
        "avoid_in_title": DEFAULT_AVOID_TITLE,
        "pushover_user": f["pushover_user"],
        "_form": {k: f[k] for k in ("titles", "core_words", "cities", "remote_only", "level", "min_salary", "lacks", "avoid", "pushover_user")},
    }


def hash_code(code: str, salt: bytes) -> str:
    return hashlib.pbkdf2_hmac("sha256", code.strip().encode(), salt, 200_000).hex()


def new_auth() -> tuple:
    """Returns (edit_code_to_show_once, auth.json dict holding only a salted hash)."""
    code = pysecrets.token_urlsafe(9)
    salt = pysecrets.token_bytes(16)
    return code, {"salt": salt.hex(), "hash": hash_code(code, salt)}


def check_code(auth: dict, code: str) -> bool:
    try:
        return hmac.compare_digest(hash_code(code, bytes.fromhex(auth["salt"])), auth["hash"])
    except (KeyError, ValueError, TypeError):
        return False


# ---------------------------------------------------------------- GitHub (private profiles repo)
class Repo:
    def __init__(self, token: str, repo: str = ""):
        self.s = requests.Session()
        self.s.headers.update({"Authorization": f"Bearer {token}", "Accept": "application/vnd.github+json", "User-Agent": "jobwatch-signup"})
        self.repo = repo or f"{self.s.get(f'{API}/user', timeout=30).json()['login']}/{PROFILES_REPO_NAME}"

    def get(self, path: str):
        """Returns (bytes, sha) or None if the file doesn't exist."""
        r = self.s.get(f"{API}/repos/{self.repo}/contents/{path}", headers={"Accept": "application/vnd.github+json"}, timeout=30)
        if r.status_code == 404:
            return None
        r.raise_for_status()
        j = r.json()
        return base64.b64decode(j["content"]), j["sha"]

    def put(self, path: str, data: bytes, message: str) -> None:
        cur = self.get(path)
        body = {"message": message, "content": base64.b64encode(data).decode()}
        if cur:
            body["sha"] = cur[1]
        r = self.s.put(f"{API}/repos/{self.repo}/contents/{path}", json=body, timeout=30)
        r.raise_for_status()

    def get_json(self, path: str):
        got = self.get(path)
        return json.loads(got[0].decode("utf-8")) if got else None


def free_slug(repo: Repo, first_name: str) -> str:
    base = slugify(first_name)
    for _ in range(20):
        slug = f"{base}-{pysecrets.token_hex(2)}"
        if repo.get(f"profiles/{slug}/profile.json") is None:
            return slug
    raise RuntimeError("could not find a free ID")


def save_person(repo: Repo, slug: str, form: dict, resume_text: str = "", auth: dict = None, email: str = None, created: str = "") -> None:
    """Commit order matters: profile.json goes LAST. The engine ignores a folder without it, so it never sees a half-written person."""
    base = f"profiles/{slug}"
    if resume_text:
        repo.put(f"{base}/resume.txt", resume_text.encode("utf-8"), f"resume for {slug}")
    if auth:
        repo.put(f"{base}/auth.json", json.dumps(auth).encode(), f"auth for {slug}")
    if email is not None:
        repo.put(f"{base}/contact.json", json.dumps({"email": email, "created": created}).encode(), f"contact for {slug}")
    repo.put(f"{base}/profile.json", json.dumps(build_profile(form), indent=1, ensure_ascii=False).encode("utf-8"), f"profile for {slug}")


def page_message(repo: Repo, slug: str) -> str:
    """The line the person sees about their page (public.json is written back by the engine after a run)."""
    info = None
    try:
        info = repo.get_json(f"profiles/{slug}/public.json")
    except Exception:
        pass
    url = (info or {}).get("gist_url", "")
    if url.startswith(GIST_PREFIX):
        return (f"**Your page:** {url} — refreshes 7 AM and 2 PM Pacific on weekdays. First results within one business day.  \n"
                f"Last run: {(info or {}).get('last_run', '?')} · jobs in the last 7 days: {(info or {}).get('jobs_last_7d', '?')}")
    return ("Your page isn't ready yet. It is created on the next run (7 AM or 2 PM Pacific, weekdays) — first results within one business day. "
            "Come back to the *Edit my filters* tab with your ID and edit code to get the link.")


def login(repo: Repo, slug: str, code: str):
    """Returns the stored profile.json if (slug, code) are right, else None. Same answer for 'no such ID' and 'wrong code'."""
    slug = slug.strip().lower()
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,39}", slug):
        return None
    auth = repo.get_json(f"profiles/{slug}/auth.json")
    if not auth or not check_code(auth, code):
        return None
    return repo.get_json(f"profiles/{slug}/profile.json")


# ---------------------------------------------------------------- the page
def main() -> None:
    import streamlit as st

    st.set_page_config(page_title="Job Watch", page_icon="🔎")
    st.title("🔎 Job Watch")
    st.caption("Tell it what job you want once. It checks job boards every weekday and keeps a private page of good, still-open matches.")

    try:
        invite, token = st.secrets["INVITE_CODE"], st.secrets["GITHUB_TOKEN"]
        repo_name = st.secrets.get("PROFILES_REPO", "")
    except Exception:
        st.error("This page isn't set up yet (missing secrets). Tell Jainam.")
        st.stop()

    ss = st.session_state
    ss.setdefault("ok", False)
    ss.setdefault("tries", 0)
    if not ss["ok"]:
        code = st.text_input("Invite code", type="password")
        if code:
            if ss["tries"] >= 8:
                st.error("Too many tries. Reload the page and ask Jainam for the code.")
            elif hmac.compare_digest(code.encode(), str(invite).encode()):
                ss["ok"] = True
                st.rerun()
            else:
                ss["tries"] += 1
                st.error("That code isn't right.")
        st.stop()

    @st.cache_resource(show_spinner=False)
    def get_repo():
        return Repo(token, repo_name)

    try:
        repo = get_repo()
    except Exception:
        st.error("Can't reach the profile store right now. Please try again in a few minutes.")
        st.stop()

    def form_fields(d: dict, key: str, with_name: bool):
        """Shared form body. d = defaults. Returns the values (not yet validated)."""
        f = {}
        f["first_name"] = st.text_input("First name", d.get("first_name", ""), key=f"{key}_fn", max_chars=30, disabled=not with_name).strip()
        f["email"] = st.text_input("Email (kept private, only so Jainam can reach you)", d.get("email", ""), key=f"{key}_em", max_chars=120).strip() if with_name else None
        f["titles"] = parse_items(st.text_area("Job titles you want — one per line (up to 25)", "\n".join(d.get("titles", [])), key=f"{key}_ti", height=130,
                                               help="Write them the way a job board would, e.g. 'Marketing Analyst'."), MAX_TITLES, 80, r"\n")
        f["core_words"] = parse_items(st.text_area("Core words — a job must mention at least one (comma or new line)", ", ".join(d.get("core_words", [])), key=f"{key}_cw",
                                                   help="The words that define your field, e.g. SEO, Google Analytics, campaign."), MAX_CORE, 60, r"[\n,]")
        f["remote_only"] = st.checkbox("Remote only (skip hybrid and on-site jobs)", d.get("remote_only", True), key=f"{key}_ro")
        f["cities"] = parse_items(st.text_area("Cities (one per line, e.g. Austin, TX) — ignored if Remote only", "\n".join(d.get("cities", [])), key=f"{key}_ci", height=90), MAX_CITIES + 1, 60, r"\n")
        keys = list(LEVELS)
        f["level"] = st.selectbox("Level", keys, index=keys.index(d.get("level", "any")), format_func=LEVELS.get, key=f"{key}_lv")
        f["min_salary"] = int(st.number_input("Minimum yearly pay in $ (0 = don't filter)", 0, 1_000_000, int(d.get("min_salary", 0)), step=5000, key=f"{key}_ms"))
        f["lacks"] = parse_items(st.text_area("Certifications or clearances you DON'T have — jobs that require them are skipped", ", ".join(d.get("lacks", [])), key=f"{key}_lk",
                                              help="e.g. PMP, CPA, security clearance"), 30, 60, r"[\n,]")
        f["avoid"] = parse_items(st.text_area("Companies to avoid (comma or new line)", ", ".join(d.get("avoid", [])), key=f"{key}_av"), 60, 60, r"[\n,]")
        f["pushover_user"] = st.text_input("Pushover user key for phone alerts (optional)", d.get("pushover_user", ""), key=f"{key}_pk", max_chars=40).strip()
        return f

    tab_join, tab_edit, tab_csv = st.tabs(["Sign up", "Edit my filters", "Download my CSV"])

    with tab_join:
        with st.form("join"):
            f = form_fields({}, "join", with_name=True)
            up = st.file_uploader("Resume (PDF or Word .docx)", type=["pdf", "docx"], key="join_up")
            st.caption("Your resume is stored only in a private repository and used only to score jobs for you.")
            go = st.form_submit_button("Sign me up")
        if go and ss.get("signups", 0) >= 3:
            st.error("Too many sign-ups from this browser session.")
        elif go:
            errs = []
            try:
                f["resume_text"] = resume_to_text(up.name, up.getvalue()) if up else ""
            except ValueError as e:
                f["resume_text"], errs = "", [str(e)]
            errs = (errs or validate_form(f)) if up else validate_form(f)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                try:
                    slug = free_slug(repo, f["first_name"])
                    code_shown, auth = new_auth()
                    save_person(repo, slug, f, f["resume_text"], auth, f["email"], datetime.now(timezone.utc).isoformat(timespec="seconds"))
                    ss["signups"] = ss.get("signups", 0) + 1
                    st.success("You're in!")
                    st.markdown(f"**Your ID:** `{slug}`  \n**Your edit code:** `{code_shown}`")
                    st.warning("Write both down now — the edit code can't be shown again. You need them to change your filters or download your CSV.")
                    st.markdown(page_message(repo, slug))
                except Exception:
                    st.error("Something went wrong saving your profile. Nothing was lost on your side — please try again in a minute, or tell Jainam.")

    with tab_edit:
        c1, c2 = st.columns(2)
        eslug = c1.text_input("Your ID", key="e_slug")
        ecode = c2.text_input("Edit code", type="password", key="e_code")
        if st.button("Open my profile", key="e_open"):
            try:
                prof = login(repo, eslug, ecode)
            except Exception:
                prof = None
            if prof is None:
                ss.pop("edit", None)
                st.error("That ID and code don't match.")
            else:
                ss["edit"] = {"slug": eslug.strip().lower(), "code": ecode, "profile": prof}
        ed = ss.get("edit")
        if ed:
            st.markdown(page_message(repo, ed["slug"]))
            d = dict(ed["profile"].get("_form", {}), first_name=ed["profile"].get("first_name", ""))
            with st.form("edit"):
                f = form_fields(d, "edit", with_name=False)
                f["first_name"] = d["first_name"]
                up = st.file_uploader("Replace my resume (optional)", type=["pdf", "docx"], key="edit_up")
                save = st.form_submit_button("Save changes")
            if save:
                errs = []
                try:
                    f["resume_text"] = resume_to_text(up.name, up.getvalue()) if up else ""
                except ValueError as e:
                    f["resume_text"], errs = "", [str(e)]
                errs = errs or validate_form(f, resume_required=False)
                if errs:
                    for e in errs:
                        st.error(e)
                else:
                    try:
                        save_person(repo, ed["slug"], f, f["resume_text"])
                        st.success("Saved. Your next run (7 AM or 2 PM Pacific, weekdays) uses the new filters.")
                    except Exception:
                        st.error("Something went wrong saving. Please try again in a minute.")

    with tab_csv:
        c1, c2 = st.columns(2)
        cslug = c1.text_input("Your ID", key="c_slug")
        ccode = c2.text_input("Edit code", type="password", key="c_code")
        if st.button("Get my CSV", key="c_get"):
            try:
                ok = login(repo, cslug, ccode) is not None
                got = repo.get(f"profiles/{cslug.strip().lower()}/export.csv") if ok else None
            except Exception:
                ok, got = False, None
            if not ok:
                st.error("That ID and code don't match.")
            elif not got:
                st.info("No CSV yet. It is refreshed once a week after your first results.")
            else:
                st.download_button("Download my_jobs.csv", got[0], file_name="my_jobs.csv", mime="text/csv")


if __name__ == "__main__":
    main()
